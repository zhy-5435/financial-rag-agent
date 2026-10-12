"""格式路由读取器：按文件后缀分发到对应解析器，统一产出 RawDocument（抽取文本逐字保留）

- .txt / .md：原生读取（多编码回退），零第三方依赖；
- .pdf / .docx：LlamaIndex llama-index-readers-file（pypdf / docx2txt），
  延迟导入——Phase 0 空跑纯文本文档时无需已安装解析器包；PDF 抽取器选型（pypdf vs pymupdf4llm）留待 Phase 1 金标对比。
- .pptx：python-pptx 原生轻抽取（PptxReader 初始化强烧图像描述模型，与限域引入相左，见 _read_pptx 注释）。
二进制格式若文件名无中文（投放系统导出的哈希名，如 20260424170528_52kqyhc7p9.docx），
标题回退取正文首个非空短行（正式文书首行即《…规则》标题本身），见 _resolve_title。
"""
import re
from pathlib import Path

from src.data.ingestion.config import SCANNED_PAGE_MIN_CHARS, SUPPORTED_SUFFIXES
from src.data.ingestion.schemas import RawDocument

# 文本类编码回退顺序（金融文档导出常见 UTF-8 与 GBK 混杂）
_TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030")

_CJK = re.compile(r"[\u4e00-\u9fff]")


def _resolve_title(path: Path, text: str) -> str:
    """文件名含中文直接沿用；哈希/纯 ASCII 名回退取正文首个非空短行，仍不合规则保留文件名"""
    if _CJK.search(path.stem):
        return path.stem
    for ln in text.splitlines():
        ln = ln.strip().lstrip("#").strip()
        if ln:
            return ln if len(ln) <= 40 and not ln.isdigit() else path.stem
    return path.stem


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_SUFFIXES


def read_document(path: Path) -> RawDocument:
    """按后缀路由读取单文件并抽取文本，不支持的格式抛 ValueError。
    所有通道出口统一清洗文本（剔除无法入 SQLite 的孤立代理码点，见 _sanitize_text）"""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"不支持的文件格式：{suffix}（支持 {sorted(SUPPORTED_SUFFIXES)}）")
    if suffix in (".txt",):
        doc = _read_plain(path)
    elif suffix in (".md", ".markdown"):
        doc = _read_plain(path)
        doc.title = _first_heading(doc.text) or doc.title
    elif suffix == ".pptx":
        doc = _read_pptx(path)
    else:
        doc = _read_with_llama_index(path)
    doc.text = _sanitize_text(doc.text)
    return doc


def _sanitize_text(text: str) -> str:
    r"""剔除孤立代理码点：pypdf 按 UTF-16 解码数学粗体/私用区字符时可能只留下半个代理对
    （\ud800~\udfff 孤立项），它不是合法字符且 SQLite 存文本需转 UTF-8，写库 executemany
    会抛 UnicodeEncodeError（正式版实测招银理财说明书含此碎片致整库重建中断）。预检/空跑不写库
    故不暴露，仅在正式入库时触发——于抽取出口统一清洗，有效正文逐字不动"""
    try:
        text.encode("utf-8")
        return text
    except UnicodeEncodeError:
        return "".join(ch for ch in text if not "\ud800" <= ch <= "\udfff")


def _read_plain(path: Path) -> RawDocument:
    """纯文本读取：编码逐一回退，全部失败按 utf-8 replace 兜底（乱码交由 quality 检出）"""
    data = path.read_bytes()
    for enc in _TEXT_ENCODINGS:
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("utf-8", errors="replace")
    return RawDocument(source_path=path, title=path.stem, text=text.strip())


def _first_heading(markdown_text: str) -> str | None:
    """取 Markdown 首个一级标题作为文档标题，无则返回 None"""
    for line in markdown_text.splitlines():
        if line.startswith("# "):
            return line.lstrip("# ").strip()
    return None


def _read_pptx(path: Path) -> RawDocument:
    """PPTX 轻抽取（python-pptx 直读）。

    不用 LlamaIndex PptxReader：其初始化强烧 vit-gpt2 图像描述模型（torch + HF 下载），
    与新 transformers 不兼容，也与「限域引入、不拖生成依赖」相左。改为每页产出
    「## 第N页 标题」Markdown 标题 + 正文段落 + 表格转 Markdown 行，直接复用
    structuring 的 Markdown 通道（每页一个 Section，表格受整块保护）。
    """
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = [f"# {path.stem}"]
    for idx, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        blocks: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line and line != title:
                        blocks.append(line)
            if getattr(shape, "has_table", False):
                rows = [[c.text.strip() for c in r.cells] for r in shape.table.rows]
                rows = [r for r in rows if any(r)]
                if rows:
                    blocks.append("| " + " | ".join(rows[0]) + " |")
                    blocks.append("|" + "|---" * len(rows[0]) + "|")
                    blocks.extend("| " + " | ".join(r) + " |" for r in rows[1:])
        head = f"## 第{idx}页" + (f" {title}" if title else "")
        parts.append(head + "\n" + "\n".join(blocks))
    page_count = sum(1 for _ in prs.slides)
    text = "\n\n".join(parts)
    return RawDocument(source_path=path, title=_resolve_title(path, text), text=text, page_count=page_count)


def _read_with_llama_index(path: Path) -> RawDocument:
    """PDF/DOCX 走 LlamaIndex file readers，延迟导入并回填页数（扫描件判定用）"""
    try:
        from llama_index.readers.file import DocxReader, PDFReader
    except ImportError as e:
        raise ImportError(
            f"解析 {path.suffix} 需要安装 llama-index-readers-file（pip install llama-index-readers-file）"
        ) from e

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        # Phase 0 先用 pypdf（reader 包内置依赖）；每页一个 Document，metadata 含页码
        pages = PDFReader().load_data(path)
        page_texts = [d.text.strip() for d in pages]
        text = "\n".join(t for t in page_texts if t)
        page_count = len(pages)
        # 空白页计数供半扫描件判定（图片页抽不出字，仅页眉页脚级残留）
        empty_pages = sum(1 for t in page_texts if len(t) < SCANNED_PAGE_MIN_CHARS)
    else:  # .docx
        nodes = DocxReader().load_data(path)
        text = "\n\n".join(d.text for d in nodes)
        page_count = None
        empty_pages = None
    return RawDocument(source_path=path, title=_resolve_title(path, text), text=text.strip(),
                       page_count=page_count, empty_pages=empty_pages)
